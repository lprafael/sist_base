import openpyxl
import uuid
from datetime import date

EXCEL_PATH = 'CLUB SANTA TERESITA ATLETAS Y PADRES.xlsx'
ACADEMIA_ID = 'dc81e3c8-0edc-4309-afce-41f708166496'
SUCURSAL_ID = '00a5f636-4bee-4915-847e-d600f6965a65'

def to_title_case(name):
    particulas = {'de', 'del', 'la', 'las', 'los', 'do', 'da', 'y', 'e'}
    words = name.strip().split()
    res = []
    for i, w in enumerate(words):
        w_lower = w.lower()
        if i > 0 and w_lower in particulas:
            res.append(w_lower)
        else:
            res.append(w.capitalize())
    return " ".join(res)

def split_full_name(full_name):
    full_name = full_name.strip()
    words = full_name.split()
    if len(words) == 1:
        return to_title_case(words[0]), ""
    elif len(words) == 2:
        return to_title_case(words[0]), to_title_case(words[1])
    
    # Manejar partículas como 'DEL PUERTO', 'DO VALLE'
    if len(words) == 3:
        if words[1].upper() in ['DEL', 'DE', 'DO', 'DA', 'LA']:
            return to_title_case(words[0]), to_title_case(f"{words[1]} {words[2]}")
        if words[0].upper() in ['MARIA', 'ANA', 'JUAN', 'JOSE', 'LUIS', 'CARLOS']:
            return to_title_case(f"{words[0]} {words[1]}"), to_title_case(words[2])
        return to_title_case(words[0]), to_title_case(f"{words[1]} {words[2]}")
    
    if len(words) >= 4:
        if words[0].upper() in ['MARIA', 'ANA', 'JUAN', 'JOSE', 'LUIS', 'CARLOS']:
            return to_title_case(f"{words[0]} {words[1]}"), to_title_case(" ".join(words[2:]))
        return to_title_case(" ".join(words[:2])), to_title_case(" ".join(words[2:]))

def calcular_dv_ruc(ruc_base):
    try:
        ruc_str = str(ruc_base).strip().replace('.', '').replace(',', '')
        if not ruc_str.isdigit():
            return None
        k = 2
        total = 0
        for digit in reversed(ruc_str):
            total += int(digit) * k
            k += 1
            if k > 11:
                k = 2
        resto = total % 11
        if resto > 1:
            return str(11 - resto)
        else:
            return '0'
    except Exception:
        return None

def escape_sql(val):
    if val is None:
        return "NULL"
    s = str(val).replace("'", "''")
    return f"'{s}'"

def main():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb['Hoja1']

    tutores_dict = {} # (nombre_upper, ruc_num) -> tutor_info
    alumnos_list = []

    for r in range(3, sheet.max_row + 1):
        al_raw = sheet.cell(r, 2).value
        tut_raw = sheet.cell(r, 3).value
        ruc_raw = sheet.cell(r, 4).value

        if not al_raw and not tut_raw:
            continue

        al_str = str(al_raw).strip()
        tut_str = str(tut_raw).strip()
        ruc_str = str(ruc_raw).strip() if ruc_raw is not None else ""

        al_nom, al_ape = split_full_name(al_str)
        tut_nom, tut_ape = split_full_name(tut_str)

        ruc_limpio = ruc_str.replace('.', '').replace(' ', '')
        ruc_num = ""
        ruc_dv = ""
        if '-' in ruc_limpio:
            parts = ruc_limpio.split('-', 1)
            ruc_num = parts[0]
            ruc_dv = parts[1]
        elif ruc_limpio.isdigit():
            ruc_num = ruc_limpio
            ruc_dv = calcular_dv_ruc(ruc_num) or ""
        elif ruc_limpio.upper() in ['S/F', 'SF', 'SIN FACTURA', '']:
            ruc_num = ""
            ruc_dv = ""
        else:
            ruc_num = ruc_limpio
            ruc_dv = ""

        ruc_display = f"{ruc_num}-{ruc_dv}" if ruc_num and ruc_dv else ruc_num

        tutor_key = (tut_str.upper(), ruc_num)
        if tutor_key not in tutores_dict:
            tutor_id = str(uuid.uuid4())
            tutores_dict[tutor_key] = {
                'id': tutor_id,
                'nombre': tut_nom,
                'apellido': tut_ape,
                'nombre_completo': f"{tut_nom} {tut_ape}".strip(),
                'ruc_num': ruc_num,
                'ruc_dv': ruc_dv,
                'ruc_display': ruc_display,
                'raw_name': tut_str,
                'alumnos': []
            }
        
        alumno_id = str(uuid.uuid4())
        al_info = {
            'id': alumno_id,
            'nombre': al_nom,
            'apellido': al_ape,
            'nombre_completo': f"{al_nom} {al_ape}".strip(),
            'tutor_id': tutores_dict[tutor_key]['id'],
            'tutor_key': tutor_key,
            'ruc_display': ruc_display,
            'row': r
        }
        tutores_dict[tutor_key]['alumnos'].append(al_info)
        alumnos_list.append(al_info)

    print(f"=== RESUMEN DE PROCESAMIENTO ===")
    print(f"Total Alumnos leídos: {len(alumnos_list)}")
    print(f"Total Tutores únicos: {len(tutores_dict)}")

    # Generar SQL
    sql_lines = [
        "-- Importación de Atletas y Padres: Club Deportivo Santa Teresita de Liseux",
        "-- Archivo origen: CLUB SANTA TERESITA ATLETAS Y PADRES.xlsx",
        "BEGIN;",
        ""
    ]

    # 1. Insertar Tutores
    sql_lines.append("-- 1. TUTORES")
    for t in tutores_dict.values():
        sql_lines.append(
            f"INSERT INTO academias.tutores (id, academia_id, nombre, apellido, telefono, email, vinculo, es_pagador) "
            f"VALUES ('{t['id']}', '{ACADEMIA_ID}', {escape_sql(t['nombre'])}, {escape_sql(t['apellido'])}, NULL, NULL, 'Responsable', true);"
        )
    sql_lines.append("")

    # 2. Insertar Alumnos
    sql_lines.append("-- 2. ALUMNOS")
    for a in alumnos_list:
        tut = tutores_dict[a['tutor_key']]
        notas = f"Responsable: {tut['nombre_completo']}"
        if tut['ruc_display']:
            notas += f" | RUC: {tut['ruc_display']}"
        
        sql_lines.append(
            f"INSERT INTO academias.alumnos (id, academia_id, sucursal_id, nombre, apellido, fecha_nacimiento, foto_perfil, tipo_sangre, alergias, condiciones_medicas, seguro_medico, contacto_emergencia, estado, notas) "
            f"VALUES ('{a['id']}', '{ACADEMIA_ID}', '{SUCURSAL_ID}', {escape_sql(a['nombre'])}, {escape_sql(a['apellido'])}, NULL, NULL, NULL, NULL, NULL, NULL, NULL, 'activo', {escape_sql(notas)});"
        )
    sql_lines.append("")

    # 3. Insertar Alumno_Tutores
    sql_lines.append("-- 3. ALUMNO_TUTORES")
    for a in alumnos_list:
        sql_lines.append(
            f"INSERT INTO academias.alumno_tutores (alumno_id, tutor_id, es_tutor_principal) "
            f"VALUES ('{a['id']}', '{a['tutor_id']}', true) "
            f"ON CONFLICT (alumno_id, tutor_id) DO UPDATE SET es_tutor_principal = true;"
        )
    sql_lines.append("")

    # 4. Insertar Historial de Asociación
    sql_lines.append("-- 4. HISTORIAL_ASOCIACION")
    for a in alumnos_list:
        sql_lines.append(
            f"INSERT INTO academias.historial_asociacion (id, alumno_id, fecha_inicio, motivo_baja) "
            f"VALUES (uuid_generate_v4(), '{a['id']}', '2026-01-01', NULL);"
        )
    sql_lines.append("")

    # 5. Insertar Datos de Facturación (para tutores que tienen RUC o nombre)
    sql_lines.append("-- 5. DATOS DE FACTURACION")
    for t in tutores_dict.values():
        if t['ruc_num']:
            sql_lines.append(
                f"INSERT INTO facturacion.datos_facturacion (id, academia_id, tutor_id, receptor_ruc, receptor_dv, receptor_nombre, es_pagador_principal) "
                f"VALUES (uuid_generate_v4(), '{ACADEMIA_ID}', '{t['id']}', {escape_sql(t['ruc_num'])}, {escape_sql(t['ruc_dv'])}, {escape_sql(t['nombre_completo'])}, true);"
            )
        else:
            sql_lines.append(
                f"INSERT INTO facturacion.datos_facturacion (id, academia_id, tutor_id, receptor_ruc, receptor_dv, receptor_nombre, es_pagador_principal) "
                f"VALUES (uuid_generate_v4(), '{ACADEMIA_ID}', '{t['id']}', NULL, NULL, {escape_sql(t['nombre_completo'])}, true);"
            )
    sql_lines.append("")

    sql_lines.append("COMMIT;")
    sql_lines.append("")

    sql_content = "\n".join(sql_lines)
    with open("import_santa_teresita.sql", "w", encoding="utf-8") as f:
        f.write(sql_content)

    print(f"Archivo 'import_santa_teresita.sql' generado exitosamente con {len(sql_lines)} líneas.")

if __name__ == '__main__':
    main()
